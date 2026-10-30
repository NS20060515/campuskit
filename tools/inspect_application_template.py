# -*- coding: utf-8 -*-
"""读取申报书模板结构：段落（含样式）、表格、节设置，输出到文本文件便于查看。"""
import sys

from docx import Document
from docx.shared import Pt

src = sys.argv[1]
out = sys.argv[2]

d = Document(src)
lines = []

sec = d.sections[0]
lines.append("=== 页面设置 ===")
lines.append("  纸张: %.1f x %.1f cm" % (sec.page_width.cm, sec.page_height.cm))
lines.append("  页边距 上/下/左/右: %.1f / %.1f / %.1f / %.1f cm" % (
    sec.top_margin.cm, sec.bottom_margin.cm, sec.left_margin.cm, sec.right_margin.cm))
lines.append("")

lines.append("=== 正文段落（共 %d 段，含空段） ===" % len(d.paragraphs))
for i, p in enumerate(d.paragraphs):
    t = p.text
    st = p.style.name if p.style is not None else "?"
    align = p.alignment
    # 收集该段的有效字号
    sizes = []
    for r in p.runs:
        if r.font.size:
            sizes.append(r.font.size.pt)
    sz = ("%.1f" % max(sizes)) if sizes else "-"
    runs = len(p.runs)
    lines.append("[%03d] style=%-14s align=%-8s size=%-5s runs=%-2d | %s" % (
        i, st, str(align).replace("WD_ALIGN_PARAGRAPH.", "") if align else "-", sz, runs,
        t.replace("\n", "\\n") if t.strip() else "(空)"))
lines.append("")

lines.append("=== 表格（共 %d 个） ===" % len(d.tables))
for ti, tb in enumerate(d.tables):
    lines.append("--- 表 %d：%d 行 x %d 列 ---" % (ti + 1, len(tb.rows), len(tb.columns)))
    for ri, row in enumerate(tb.rows):
        cells = []
        seen = set()
        for c in row.cells:
            cid = id(c._tc)
            if cid in seen:
                cells.append("<合并>")
                continue
            seen.add(cid)
            txt = c.text.replace("\n", " ⏎ ").strip()
            cells.append(txt if txt else "(空)")
        lines.append("  r%02d | %s" % (ri, " ｜ ".join(cells)))
    lines.append("")

# 页眉页脚
for i, s in enumerate(d.sections):
    for kind, part in (("页眉", s.header), ("页脚", s.footer)):
        txt = " ⏎ ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
        if txt:
            lines.append("%s: %s" % (kind, txt))

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("written: " + out)
